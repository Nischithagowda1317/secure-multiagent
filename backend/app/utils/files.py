from __future__ import annotations

import io
import re
from pathlib import Path
from typing import BinaryIO

import fitz
import pandas as pd
from docx import Document


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx"}


def safe_filename(filename: str) -> str:
    stem = Path(filename).stem
    suffix = Path(filename).suffix.lower()
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._") or "document"
    return f"{stem}{suffix}"


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix == ".pdf":
        document = fitz.open(stream=data, filetype="pdf")
        pages = []
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(f"[Page {page_number}]\n{text}")
        return "\n\n".join(pages)

    if suffix == ".docx":
        document = Document(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    paragraphs.append(" | ".join(values))
        return "\n".join(paragraphs)

    if suffix == ".txt":
        return data.decode("utf-8", errors="replace")

    if suffix == ".csv":
        frame = pd.read_csv(io.BytesIO(data))
        return _frame_to_text(frame)

    frame = pd.read_excel(io.BytesIO(data))
    return _frame_to_text(frame)


def _frame_to_text(frame: pd.DataFrame, max_rows: int = 300) -> str:
    rows = [f"Columns: {', '.join(map(str, frame.columns))}"]
    sample = frame.head(max_rows).fillna("")
    for index, row in sample.iterrows():
        parts = [f"{column}={row[column]}" for column in sample.columns]
        rows.append(f"Row {index + 1}: " + "; ".join(parts))
    if len(frame) > max_rows:
        rows.append(f"[Only the first {max_rows} of {len(frame)} rows were converted to text.]")
    return "\n".join(rows)


def chunk_text(text: str, chunk_size: int = 900, overlap: int = 120) -> list[str]:
    words = text.split()
    if not words:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        start = max(start + 1, end - overlap)
    return chunks
